#!/usr/bin/env python3
"""
One-off cleanup + finalize for Cohort 2 applicant profiles.

1. Dedupe participants by (cohort_id, lower(email)) — a seed retry created duplicates.
2. Import all profiles from the 'Accepted' sheet (deduped by email).
3. Link each application_profile to its participant by email.

Run:
  DSN=postgresql://... XLSX="/path/Accepted_Rejected List.xlsx" \
  python3 scripts/finalize_cohort2_profiles.py
"""
import os
import openpyxl
import psycopg2
from psycopg2.extras import execute_values

COHORT_ID = "b528398e-985f-49e1-80e7-b1cb7f86c7c8"

HEADER_MAP = {
    "first name": "first_name",
    "last name": "last_name",
    "email address": "email",
    "phone number (whatsapp preferred)": "phone",
    "gender": "gender",
    "age range": "age_range",
    "institution / organization": "institution",
    "level of study": "level_of_study",
    "how would you describe your background": "background",
    "have you ever built or worked on a digital product before": "built_product_before",
    "what best describes your interest in this bootcamp?": "bootcamp_interest",
    "do you currently have a n idea you'd like to turn into a product?": "has_idea",
    "describe the idea you want to work on during the bootcamp": "idea_description",
    "who do you think would benetif the most from your solution?": "beneficiary",
    "what skills or knowledge are u most hoping to gain from this program?": "skills_hoping_to_gain",
    "in what ways would you be interested in contributing after the program?": "contribution_interest",
    "are you interested in being considered as a community lead  for this cohort?": "interested_community_lead",
    "why are you interested in supporting community coordination and peer engagement?": "community_lead_reason",
    "do you have any prior volunteering, leadership, or community management experience?": "prior_experience",
    "if yes, briefly describe your experience and your role": "prior_experience_detail",
    "what strengths would you bring to a learning community like this?": "community_strengths",
    "enter your scholarship code": "scholarship_code",
    "is there anything else you'd like us to know about you?": "anything_else",
    "how did you hear about us?": "heard_about_us",
    "do you have access to laptop/desktop computer?": "has_laptop",
    "do you have stable internet connection?": "has_internet",
    "what do you think makes you a good for this bootcamp?": "good_fit_reason",
}

COLUMNS = [
    "cohort_id", "decision", "first_name", "last_name", "email", "phone", "gender",
    "age_range", "institution", "level_of_study", "background", "built_product_before",
    "bootcamp_interest", "has_idea", "idea_description", "beneficiary", "skills_hoping_to_gain",
    "giving_back_importance", "contribution_interest", "interested_community_lead",
    "community_lead_reason", "prior_experience", "prior_experience_detail", "community_strengths",
    "scholarship_code", "can_commit", "anything_else", "heard_about_us", "has_laptop",
    "has_internet", "good_fit_reason",
]


def map_header(raw):
    if not raw:
        return None
    key = " ".join(str(raw).strip().lower().split())
    if key in HEADER_MAP:
        return HEADER_MAP[key]
    if key.startswith("tnx solve focuses") or "why is giving back" in key:
        return "giving_back_importance"
    if key.startswith("the program requires weekly"):
        return "can_commit"
    return None


def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def main():
    dsn = os.environ["DSN"]
    xlsx = os.environ["XLSX"]
    c = psycopg2.connect(dsn, connect_timeout=10)
    c.autocommit = True
    cur = c.cursor()

    # 1. dedupe participants
    cur.execute(
        """
        delete from public.participants p using (
          select id, row_number() over (partition by cohort_id, lower(email) order by created_at) rn
          from public.participants where cohort_id=%s
        ) d
        where p.id=d.id and d.rn>1
        """,
        (COHORT_ID,),
    )
    print("deleted dup participants:", cur.rowcount)
    cur.execute("select count(*) from public.participants where cohort_id=%s", (COHORT_ID,))
    print("participants now:", cur.fetchone()[0])

    # 2. parse accepted profiles (dedupe by email)
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = [w for w in wb.worksheets if w.title.strip().lower() == "accepted"][0]
    header = None
    by_email = {}
    for row in ws.iter_rows(values_only=True):
        if header is None:
            if row and any(row):
                header = [map_header(x) for x in row]
            continue
        if not row or not any(row):
            continue
        rec = {"cohort_id": COHORT_ID, "decision": "Accepted"}
        for ci, ck in enumerate(header):
            if ck and ci < len(row):
                rec[ck] = clean(row[ci])
        email = (rec.get("email") or "").strip().lower()
        if not email:
            continue
        rec["email"] = email
        by_email[email] = rec  # last one wins on dupes

    values = [[rec.get(col) for col in COLUMNS] for rec in by_email.values()]
    placeholders = "(" + ",".join(["%s"] * len(COLUMNS)) + ")"
    updates = ",".join([f"{col}=excluded.{col}" for col in COLUMNS if col != "email"])
    sql = (
        f"insert into public.application_profiles ({','.join(COLUMNS)}) values %s "
        f"on conflict (email) do update set {updates}"
    )
    execute_values(cur, sql, values, template=placeholders)
    print("upserted profiles:", len(values))

    # 3. link profiles -> participants by email
    cur.execute(
        """
        update public.application_profiles ap
        set participant_id = p.id
        from public.participants p
        where p.cohort_id = %s and lower(p.email) = ap.email
        """,
        (COHORT_ID,),
    )
    print("linked profiles:", cur.rowcount)

    cur.execute("select count(*) from public.application_profiles")
    print("total profiles:", cur.fetchone()[0])
    c.close()


if __name__ == "__main__":
    main()
